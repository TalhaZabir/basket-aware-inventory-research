from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "sections_4_5_6_methodology_results_discussion.md"
OUTPUT = ROOT / "sections_4_5_6_methodology_results_discussion.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "")
    text = text.replace("\\&", "&")
    return text


def equation_text(lines: list[str]) -> list[str]:
    joined = " ".join(line.strip() for line in lines)
    if "text{retain line}" in joined:
        return [
            "Retain line = 1 if:",
            "invoice, stock code, description, and timestamp are present; invoice does not begin with C; quantity is greater than 0; unit price is greater than 0; and stock code begins with a numeric character.",
        ]
    joined = joined.replace("\\mathbf{", "").replace("}", "")
    joined = re.sub(r"\\text\{([^}]*)\}", r"\1", joined)
    joined = joined.replace("\\max", "max").replace("\\mathbf", "")
    joined = joined.replace("\\sum", "sum").replace("\\frac", "ratio")
    joined = joined.replace("\\mathbf{1}", "1").replace("\\mathbb", "")
    joined = joined.replace("\\", "")
    joined = re.sub(r"\s+", " ", joined).strip()
    return [joined]


def set_default_font(doc: Document) -> None:
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    title = doc.styles["Title"]
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    subtitle = doc.styles["Subtitle"]
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = RGBColor(89, 89, 89)
    subtitle.paragraph_format.space_after = Pt(18)
    for name, size, color in [("Heading 1", 15, "000000"), ("Heading 2", 12, "1F4E79"), ("Heading 3", 10.5, "365F91")]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    repeat_table_header(header)
    for index in range(cols):
        cell = header.cells[index]
        cell.text = clean_inline(rows[0][index] if index < len(rows[0]) else "")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E79")
        set_cell_borders(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Arial"
                run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
                run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows[1:], start=1):
        row = table.add_row()
        prevent_row_split(row)
        for col_index in range(cols):
            cell = row.cells[col_index]
            cell.text = clean_inline(values[col_index] if col_index < len(values) else "")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F2F6FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 or any(ch.isalpha() for ch in cell.text) else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
                    run.font.size = Pt(8.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    i = 0
    first_h1 = True
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        figure_specs = {
            "[[FRONTIER_FIGURE]]": (
                "frontier_plot.png",
                "Figure 3. Service versus average inventory frontier across representative scenarios; the primary alpha is selected by matched-inventory validation.",
                "Service versus average inventory frontier for independent and basket-aware policies across representative scenarios",
                "Service versus average inventory frontier",
            ),
            "[[COST_PENALTY_FIGURE]]": (
                "cost_vs_penalty.png",
                "Figure 4. Cost differences as the shortage penalty changes; crossings follow arithmetically from the frozen cost formula and break-even p/h ratios.",
                "Simulated policy cost differences as the shortage penalty varies",
                "Cost versus shortage penalty",
            ),
            "[[BASKET_SIZE_FIGURE]]": (
                "basket_size_distribution.png",
                "Figure 1. Distribution of retained transaction-line basket sizes in the cleaned UCI data.",
                "Distribution of retained basket sizes measured in transaction lines",
                "Basket-size distribution",
            ),
            "[[GRAPH_SUMMARY_FIGURE]]": (
                "graph_summary.png",
                "Figure 2. Training-only co-purchase graph summary for the executed product universe.",
                "Training-only co-purchase graph degree and edge-weight summary",
                "Co-purchase graph summary",
            ),
            "[[FOREST_FIGURE]]": (
                "forest_plot.png",
                "Figure 5. Forest plot of PA1 minus C4 and PA1 minus C1 item-fill differences across all twelve synthetic scenarios; whiskers are 95% replication intervals.",
                "Scenario forest plot for matched item-fill differences",
                "Scenario-level matched contrasts",
            ),
            "[[STRATA_FIGURE]]": (
                "mechanism_strata.png",
                "Figure 6. PA1 minus C4 item-fill differences in low-popularity/high-lift and high-popularity/high-lift product strata; whiskers are 95% replication intervals.",
                "Graph signal by popularity and lift stratum",
                "Graph signal by popularity stratum",
            ),
            "[[REPLICATION_FIGURE]]": (
                "dataset_replication.png",
                "Figure 7. PA1 minus C4 item-fill differences in UCI Online Retail II and the tightened dunnhumby Complete Journey replay; the external replay is a stress check, not a causal validation.",
                "Replication of the item-fill contrast across basket datasets",
                "Replication across basket datasets",
            ),
        }
        if line in figure_specs:
            filename, caption_text, alt_text, title_text = figure_specs[line]
            figure_path = ROOT / "research_data_package" / "results" / filename
            if figure_path.exists():
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                inline_shape = paragraph.add_run().add_picture(str(figure_path), width=Inches(6.7))
                inline_shape._inline.docPr.set("descr", alt_text)
                inline_shape._inline.docPr.set("title", title_text)
                caption = doc.add_paragraph(caption_text)
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in caption.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)
                    run.font.italic = True
            else:
                doc.add_paragraph(f"{caption_text} Plot unavailable; see `results/{filename}` in the research package.")
            i += 1
            continue
        if line.startswith("# "):
            if not first_h1:
                doc.add_page_break()
            first_h1 = False
            doc.add_heading(clean_inline(line[2:]), level=1)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=2)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=3)
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                parsed.append(cells)
            add_table(doc, parsed)
            continue
        if line.startswith("\\["):
            eq_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("\\]"):
                eq_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            for eq in equation_text(eq_lines):
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(8)
                run = paragraph.add_run(eq)
                run.font.name = "Cambria"
                run._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
                run._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
                run.font.size = Pt(10)
                run.font.italic = True
            continue
        if line.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                paragraph = doc.add_paragraph(style="List Bullet")
                paragraph.add_run(clean_inline(lines[i].strip()[2:]))
                i += 1
            continue
        if re.match(r"^\d+\. ", line):
            while i < len(lines) and re.match(r"^\d+\. ", lines[i].strip()):
                paragraph = doc.add_paragraph(style="List Number")
                paragraph.add_run(clean_inline(re.sub(r"^\d+\. ", "", lines[i].strip())))
                i += 1
            continue
        paragraph = doc.add_paragraph(clean_inline(line))
        i += 1


def build() -> None:
    doc = Document()
    set_default_font(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    r = footer.add_run("Basket Aware Inventory Research | Sections 4–6 | Page ")
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.add_run("Basket Aware Inventory Research")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Research Methodology Results and Discussion")
    intro = doc.add_paragraph()
    intro.add_run("Scope. ").bold = True
    intro.add_run("This document presents the executed methodology, results, and discussion for the basket-aware inventory study. Observed transaction and forecast results are separated from simulation outcomes under synthetic inventory assumptions, and no historical inventory or retailer-savings claims are made.")
    doc.add_paragraph()
    parse_markdown(doc, SOURCE.read_text(encoding="utf-8"))

    doc.core_properties.title = "Basket Aware Inventory Research Methodology Results and Discussion"
    doc.core_properties.subject = "Sections 4 5 and 6 of the basket-aware inventory research study"
    doc.core_properties.author = "Research Project"
    doc.core_properties.comments = "Prepared from the project research data package and literature review."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
